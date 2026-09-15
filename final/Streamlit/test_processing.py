import tempfile
import unittest
from pathlib import Path
from contextlib import closing
from unittest.mock import patch
import database
import processing as p
import management as m
from rag_service import DocumentRAG


class ProcessingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.db = self.root / 'dms_system.db'
        database.init_db(self.db)

    def document(self, name, data):
        return m.save_document(1, self.root, name, '', '', None, [], upload=(name,data), path=self.db)

    def test_text_csv_and_rag(self):
        doc = self.document('table.csv', 'Нэр;Цалин\n"Бат; Дорж";500\n'.encode())
        state, message, sections = p.process_document(doc,self.root,path=self.db)
        self.assertEqual(state,'ready')
        self.assertIn('Бат; Дорж', sections[1][1])
        self.assertEqual(len(sections),2)
        with closing(database.open_database(self.db)) as conn:
            stored = conn.execute('SELECT file_path FROM documents WHERE id=?',(doc,)).fetchone()[0]
        rag = DocumentRAG([dict(id=doc,title='Цалин',file_path=stored)],self.root)
        self.assertEqual(rag.document_count,1)
        self.assertTrue(any('500' in c.text for c in rag.chunks))

    def test_word_heading_and_table_order(self):
        from docx import Document
        doc = Document()
        doc.add_heading('Policy',level=1)
        doc.add_paragraph('Before')
        doc.add_table(rows=1,cols=2).rows[0].cells[0].text = 'Table content'
        doc.add_paragraph('After')
        source = self.root/'policy.docx'
        doc.save(source)
        state, _, sections = p.extract(source)
        self.assertEqual(state,'ready')
        self.assertIn('Before',sections[1][1])
        self.assertIn('Table content',sections[2][1])
        self.assertIn('After',sections[3][1])
        self.assertIn('Policy',sections[2][0])

    def test_excel_multiple_sheets_and_formula_warning(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Цалин'
        sheet.append(['Ажилтан','Дүн'])
        sheet.append(['Бат',500])
        sheet.append(['Нийт','=SUM(B2:B2)'])
        workbook.create_sheet('Other').append(['Hello'])
        source = self.root/'book.xlsx'
        workbook.save(source)
        workbook.close()
        state, warning, sections = p.extract(source)
        self.assertEqual(state,'partial')
        self.assertIn('томьёо',warning)
        self.assertIn('Цалин',sections[1][0])
        self.assertIn('500',sections[1][1])
        self.assertIn('Other',sections[-1][0])

    def test_pdf_page_labels_and_empty_page_warning(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=300,height=300)
        font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 20 200 Td (Policy text) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.add_blank_page(width=300,height=300)
        source = self.root/'policy.pdf'
        writer.write(source)
        state, warning, sections = p.extract(source)
        self.assertEqual(state,'partial')
        self.assertEqual(sections[0][0],'1-р хуудас')
        self.assertIn('Policy text',sections[0][1])
        self.assertIn('2',warning)

    def test_cache_replacement_restoration_and_failures(self):
        doc = self.document('policy.txt',b'original')
        self.assertEqual(p.process_document(doc,self.root,path=self.db)[2][0][1],'original')
        with patch.object(p,'extract',side_effect=AssertionError('should use cache')):
            self.assertEqual(p.process_document(doc,self.root,path=self.db)[0],'ready')
        m.save_document(1,self.root,'Policy','','',None,[],upload=('policy.txt',b'changed'),document_id=doc,path=self.db)
        self.assertEqual(p.process_document(doc,self.root,path=self.db)[2][0][1],'changed')
        m.restore_version(1,doc,1,self.root,path=self.db)
        self.assertEqual(p.process_document(doc,self.root,path=self.db)[2][0][1],'original')
        with patch.object(p,'extract',side_effect=ValueError('broken')):
            result = p.process_document(doc,self.root,path=self.db,force=True)
        self.assertEqual(result[0],'error')
        self.assertEqual(result[2],[])
        self.assertEqual(p.process_document(doc,self.root,path=self.db,force=True)[0],'ready')

    def test_stale_worker_cannot_publish_after_deletion(self):
        doc = self.document('policy.txt',b'original')
        def delete_during_extract(source):
            m.change_document_status(1,doc,path=self.db)
            return 'ready','',[('line','old')]
        with patch.object(p,'extract',side_effect=delete_during_extract):
            result = p.process_document(doc,self.root,path=self.db)
        self.assertEqual(result[2],[])

    def test_images_require_ocr_and_missing_file_errors(self):
        from PIL import Image
        source = self.root/'scan.png'
        Image.new('RGB',(10,10),'white').save(source)
        self.assertEqual(p.extract(source)[0],'needs_ocr')
        doc = self.document('policy.txt',b'original')
        with closing(database.open_database(self.db)) as conn, conn:
            conn.execute("UPDATE documents SET file_path='missing.txt' WHERE id=?",(doc,))
        self.assertEqual(p.process_document(doc,self.root,path=self.db)[0],'error')


if __name__ == '__main__':
    unittest.main()
