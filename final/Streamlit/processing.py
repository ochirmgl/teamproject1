"""Local, version-aware extraction. No provider requests or OCR are performed."""
import csv
import hashlib
import json
from io import StringIO
from pathlib import Path
from contextlib import closing
from uuid import uuid4

from database import open_database

LABELS = {'pending':'Хүлээгдэж байна', 'processing':'Боловсруулж байна',
          'ready':'Бэлэн', 'partial':'Хэсэгчлэн бэлэн', 'needs_ocr':'OCR шаардлагатай',
          'error':'Алдаа', 'unsupported':'Дэмжигдээгүй'}
MAX_CHARS = 5_000_000


def resolve_file(root, stored):
    candidate = Path(str(stored).replace('\\', '/'))
    if not candidate.is_absolute():
        candidate = Path(root) / candidate
    return candidate


def extract(path):
    """Return (status, warning, [(source location, text)]), preserving structure."""
    path = Path(path)
    if path.stat().st_size > 50 * 1024 * 1024:
        raise ValueError('Файл 50 MB-аас том байна.')
    ext = path.suffix.lower()
    sections, warnings = [], []
    total = 0

    def add(label, text):
        nonlocal total
        text = text.replace('\x00', '').strip()
        if text:
            total += len(text)
            if total > MAX_CHARS:
                raise ValueError('Задалсан агуулга хэт том байна. Файлыг жижиг хэсгүүдэд хуваана уу.')
            sections.append((label, text))

    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('Нууц үгтэй PDF. Нууцлалгүй хуулбар оруулна уу.')
        empty = []
        for number, page in enumerate(reader.pages, 1):
            text = (page.extract_text(extraction_mode='layout') or '') if '/Contents' in page else ''
            add(f'{number}-р хуудас', text)
            if not text.strip():
                empty.append(str(number))
        if empty:
            warnings.append('Текстгүй хуудас: ' + ', '.join(empty) + '. Скан бол OCR шаардлагатай.')
        if not sections:
            return 'needs_ocr', 'PDF-д уншигдах текст алга. Скан бол OCR шаардлагатай.', []
    elif ext == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        document = Document(path)
        heading, paragraph_number, table_number = 'Word баримт', 0, 0
        for element in document.element.body.iterchildren():
            if element.tag.endswith('}p'):
                paragraph = Paragraph(element, document)
                paragraph_number += 1
                if paragraph.style and paragraph.style.name.startswith('Heading'):
                    heading = paragraph.text.strip() or heading
                add(f'{heading} · догол мөр {paragraph_number}', paragraph.text)
            elif element.tag.endswith('}tbl'):
                table_number += 1
                table = Table(element, document)
                for number, row in enumerate(table.rows, 1):
                    add(f'{heading} · хүснэгт {table_number}, мөр {number}',
                        ' | '.join(cell.text for cell in row.cells))
        # Word page numbers depend on layout; do not invent them.
    elif ext in {'.txt', '.csv'}:
        text = path.read_text(encoding='utf-8-sig')
        if ext == '.txt':
            lines = text.splitlines()
            for start in range(0, len(lines), 40):
                add(f'Мөр {start+1}–{min(start+40,len(lines))}', '\n'.join(lines[start:start+40]))
        else:
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            header = None
            for number, row in enumerate(csv.reader(StringIO(text), dialect), 1):
                if header is None:
                    header = row
                add(f'CSV бичлэг {number}', ' | '.join(f'{header[i] if i < len(header) and header[i] else i+1}: {cell}' for i, cell in enumerate(row)))
    elif ext in {'.xlsx', '.xls'}:
        if ext == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            cached = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in workbook:
                    if sheet.max_row > 100000 or sheet.max_column > 1000:
                        raise ValueError('Хүснэгт хэт том байна. Жижиг файл болгон хуваана уу.')
                    for number, (row, values) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
                        parts = []
                        for cell, value in zip(row, values):
                            if cell.value is None:
                                continue
                            content = str(cell.value)
                            if cell.data_type == 'f':
                                if value.value is None:
                                    content += ' [тооцоолсон утга хадгалагдаагүй]'
                                    warnings.append('Зарим томьёоны тооцоолсон утга хадгалагдаагүй. Excel-д тооцоолж хадгална уу.')
                                else:
                                    content += f' = {value.value}'
                            parts.append(f'{cell.column_letter}: {content}')
                        add(f'{sheet.title} · мөр {number}', ' | '.join(parts))
            finally:
                workbook.close()
                cached.close()
        else:
            try:
                import xlrd
            except ImportError as exc:
                raise ValueError('XLS уншихад xlrd шаардлагатай. Эсвэл XLSX болгон хадгална уу.') from exc
            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                for sheet in workbook.sheets():
                    if sheet.nrows > 100000:
                        raise ValueError('Хүснэгт хэт том байна.')
                    for number in range(sheet.nrows):
                        add(f'{sheet.name} · мөр {number+1}', ' | '.join(str(v) for v in sheet.row_values(number)))
            finally:
                workbook.release_resources()
    elif ext in {'.png','.jpg','.jpeg','.webp'}:
        return 'needs_ocr', 'Зураг уншихад OCR шаардлагатай. OCR одоогоор холбогдоогүй.', []
    else:
        return 'unsupported', 'Энэ төрлийг задлах боломжгүй. DOC бол DOCX болгон хадгална уу.', []
    if not sections:
        return 'error', 'Уншигдах агуулга олдсонгүй.', []
    return ('partial' if warnings else 'ready'), '\n'.join(dict.fromkeys(warnings)), sections


def invalidate(conn, document_id, stored):
    conn.execute("""INSERT INTO document_processing(document_id,file_path) VALUES (?,?)
        ON CONFLICT(document_id) DO UPDATE SET file_path=excluded.file_path,status='pending',
        fingerprint='',message='',sections_json='[]',run_id='',updated_at=CURRENT_TIMESTAMP""", (document_id,stored))


def process_document(document_id, root, path=None, force=False):
    """Persist status and content; reject stale worker results after replacement/deletion."""
    with closing(open_database(path)) as conn:
        row = conn.execute("SELECT file_path FROM documents WHERE id=? AND status='active'", (document_id,)).fetchone()
        if not row:
            return 'error', 'Баримт идэвхгүй байна.', []
        stored = row[0]
        source = resolve_file(root, stored)
        fingerprint = ''
        try:
            if source.stat().st_size > 50 * 1024 * 1024:
                raise ValueError('Файл 50 MB-аас том байна.')
            fingerprint = hashlib.sha256(source.read_bytes()).hexdigest()
            previous = conn.execute('SELECT status,message,sections_json,fingerprint,file_path FROM document_processing WHERE document_id=?', (document_id,)).fetchone()
            if not force and previous and previous[3:] == (fingerprint, stored) and previous[0] not in ('pending','processing'):
                return previous[0], previous[1], json.loads(previous[2])
        except (OSError, ValueError) as exc:
            failure = str(exc)
        else:
            failure = None
        token = uuid4().hex
        conn.execute("""INSERT INTO document_processing(document_id,file_path,status,run_id) VALUES (?,?,'processing',?)
            ON CONFLICT(document_id) DO UPDATE SET file_path=excluded.file_path,status='processing',
            message='',sections_json='[]',run_id=excluded.run_id,updated_at=CURRENT_TIMESTAMP""", (document_id,stored,token))
        conn.commit()
        try:
            if failure:
                raise ValueError(failure)
            status, message, sections = extract(source)
            if hashlib.sha256(source.read_bytes()).hexdigest() != fingerprint:
                raise ValueError('Файл боловсруулалтын үед өөрчлөгдсөн. Дахин боловсруулна уу.')
        except Exception as exc:
            status, message, sections = 'error', f'Уншиж чадсангүй: {exc}', []
        result = conn.execute("""UPDATE document_processing SET status=?,message=?,sections_json=?,fingerprint=?,updated_at=CURRENT_TIMESTAMP
            WHERE document_id=? AND run_id=? AND EXISTS
            (SELECT 1 FROM documents WHERE id=? AND file_path=? AND status='active')""",
            (status,message,json.dumps(sections,ensure_ascii=False),fingerprint,document_id,token,document_id,stored))
        conn.commit()
        if not result.rowcount:
            conn.execute("UPDATE document_processing SET status='pending',sections_json='[]' WHERE document_id=? AND run_id=?", (document_id,token))
            conn.commit()
            return 'pending', 'Баримт өөрчлөгдсөн тул дахин боловсруулна уу.', []
        return status, message, sections


if __name__ == '__main__':
    from collections import Counter
    from database import init_db, DB_PATH
    init_db()
    with closing(open_database()) as connection:
        ids = [r[0] for r in connection.execute("SELECT id FROM documents WHERE status='active'")]
    results = [(doc, process_document(doc, DB_PATH.parent)) for doc in ids]
    print(dict(Counter(result[0] for _, result in results)))
    for doc, result in results:
        if result[1]:
            print(f'Document {doc}: {result[1]}')
